#!/usr/bin/env python3
"""
Staffing Sheet Generator
Converts HR export files into a formatted staffing sheet for 17x11 print.
"""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import sys
import os
import argparse


# Department mapping to sections
DEPARTMENT_MAPPING = {
    # Office section
    'EXEC': {'section': 'Office', 'subsection': 'Executives'},
    'OFFICE': {'section': 'Office', 'subsection': 'Office'},
    'HR': {'section': 'Office', 'subsection': 'HR'},
    'CS': {'section': 'Office', 'subsection': 'Cust Service'},
    'ESTIMATIONS': {'section': 'Office', 'subsection': 'Estimating'},
    'FINC': {'section': 'Office', 'subsection': 'Accounting'},
    'ENGR': {'section': 'Office', 'subsection': 'Engineering'},
    'QUALITY': {'section': 'Office', 'subsection': 'Quality'},
    'SHIP': {'section': 'Office', 'subsection': 'Pack/Ship'},
    'INVENTORY': {'section': 'Office', 'subsection': 'Inventory'},
    
    # Fabrication section
    'CUT': {'section': 'Fabrication', 'subsection': 'Cutting'},
    'BEND': {'section': 'Fabrication', 'subsection': 'Bending'},
    'PCO': {'section': 'Fabrication', 'subsection': 'Welding Manager'},
    'WLDSM': {'section': 'Fabrication', 'subsection': 'Small Weld'},
    'WLDMD': {'section': 'Fabrication', 'subsection': 'Med Weld'},
    'WLDLG': {'section': 'Fabrication', 'subsection': 'Large Weld'},
    'WLDSS': {'section': 'Fabrication', 'subsection': 'SS Weld'},
    'MAINT': {'section': 'Fabrication', 'subsection': 'Maintenance'},
    
    # Finishing section
    'PAINT': {'section': 'Finishing', 'subsection': 'Paint'},
    'ASMBL': {'section': 'Finishing', 'subsection': 'Assembly'},
    'LPO': {'section': 'Finishing', 'subsection': 'MFG Engineering'},
    'GPO': {'section': 'Finishing', 'subsection': 'Float'},
}

# Section order for layout
SECTION_ORDER = ['Office', 'Fabrication', 'Finishing']

# Subsection order within each section
SUBSECTION_ORDER = {
    'Office': ['Executives', 'HR', 'Cust Service', 'Estimating', 'Office', 
               'Engineering', 'Quality', 'Pack/Ship', 'Inventory', 'Accounting', 
               'Purchasing', 'Sales'],
    'Fabrication': ['Cutting', 'Bending', 'Welding Manager', 'Small Weld', 
                    'Med Weld', 'Large Weld', 'SS Weld', 'Maintenance'],
    'Finishing': ['Finishing Manager', 'Paint', 'Assembly', 'MFG Engineering', 
                  'Float', 'Systems', 'Scheduling', 'Operations Manager']
}

# Subsections that use simplified layout (Name and Title only, no Shift/ID/Codes)
SIMPLE_LAYOUT_SUBSECTIONS = ['Executives', 'HR']


def find_header_row(file_path):
    """Find the row containing 'Default Department' in the file."""
    try:
        df_raw = pd.read_excel(file_path, header=None)
        for i, row in df_raw.iterrows():
            if row[0] == "Default Department":
                return i
    except Exception:
        pass
    return None


def read_employee_file(file_path):
    """Read employee data from Excel or CSV file."""
    print(f"Reading file: {file_path}")
    
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
    
    # Determine file type and read accordingly
    file_ext = os.path.splitext(file_path)[1].lower()
    
    if file_ext in ['.xls', '.xlsx']:
        # Find header row for .xls files that may have extra header rows
        header_row = find_header_row(file_path)
        if header_row is not None:
            df = pd.read_excel(file_path, skiprows=header_row, header=0)
        else:
            df = pd.read_excel(file_path)
    elif file_ext == '.csv':
        df = pd.read_csv(file_path)
    else:
        raise ValueError(f"Unsupported file format: {file_ext}")
    
    print(f"  Loaded {len(df)} rows")
    return df


def merge_employee_data(df1, df2):
    """Merge and deduplicate employee data from two DataFrames."""
    print("Merging employee data...")
    
    # Combine both dataframes, putting df2 first to prioritize its data
    combined = pd.concat([df2, df1], ignore_index=True)
    
    # Remove duplicates based on Employee Id, keeping the first occurrence (from df2)
    merged = combined.drop_duplicates(subset=['Employee Id'], keep='first')
    
    print(f"  Combined {len(df2)} + {len(df1)} = {len(combined)} rows")
    print(f"  After deduplication: {len(merged)} rows")
    
    return merged


def filter_active_employees(df):
    """Filter to show only active employees."""
    print("Filtering to active employees...")
    
    # Filter out terminated employees
    # Keep "Active" and "Not In Payroll" (for temps/trainees)
    active = df[df['Employee Status'].isin(['Active', 'Not In Payroll'])].copy()
    
    print(f"  Kept {len(active)} active employees from {len(df)} total")
    return active


def extract_status_codes(row):
    """Extract status codes from Primary and Secondary Status columns."""
    codes = []
    
    # L = Line lead (Primary Status)
    if pd.notna(row['Primary Status ']) and 'Line lead' in str(row['Primary Status ']):
        codes.append('L')
    
    # O = Offsite employee (Primary Status)
    if pd.notna(row['Primary Status ']) and 'Offsite' in str(row['Primary Status ']):
        codes.append('O')
    
    # S = Safety Team (Secondary Status)
    if pd.notna(row['Secondary Status ']) and 'Safety Team' in str(row['Secondary Status ']):
        codes.append('S')
    
    # P = Part Time (Secondary Status)
    if pd.notna(row['Secondary Status ']) and 'Part Time' in str(row['Secondary Status ']):
        codes.append('P')
    
    # T = Trainee/Temp (Not In Payroll status)
    if row['Employee Status'] == 'Not In Payroll':
        codes.append('T')
    
    # Return codes with spaces for formatting (e.g., " L S")
    if codes:
        return ' ' + ' '.join(codes)
    return '  '


def extract_shift(work_schedule):
    """Extract shift number from work schedule."""
    if pd.isna(work_schedule):
        return '1'
    
    schedule = str(work_schedule).lower()
    if '1st' in schedule:
        return '1'
    elif '2nd' in schedule:
        return '2'
    elif '3rd' in schedule:
        return '3'
    elif 'part' in schedule:
        return 'P'
    else:
        return '1'  # Default to 1 for office hours


def simplify_job_title(job_title):
    """Simplify job title by removing department prefix."""
    if pd.isna(job_title):
        return ''
    
    title = str(job_title)
    # Remove department prefix (e.g., "WLDMD - Paint Prep" -> "Paint Prep")
    if ' - ' in title:
        return title.split(' - ', 1)[1]
    return title


def map_department(dept_code):
    """Map department code to section and subsection."""
    if pd.isna(dept_code):
        return {'section': 'Other', 'subsection': 'Other'}
    
    dept = str(dept_code).upper().strip()
    return DEPARTMENT_MAPPING.get(dept, {'section': 'Other', 'subsection': 'Other'})


def process_employee_data(df):
    """Process employee data: extract codes, map departments, etc."""
    print("Processing employee data...")
    
    # Add computed columns
    df['Status Codes'] = df.apply(extract_status_codes, axis=1)
    df['Shift'] = df['Work Schedule'].apply(extract_shift)
    df['Simple Title'] = df['Jobs (HR)(1)'].apply(simplify_job_title)
    
    # Map departments
    df['Section'] = df['Default Department'].apply(lambda x: map_department(x)['section'])
    df['Subsection'] = df['Default Department'].apply(lambda x: map_department(x)['subsection'])
    
    # Sort by section, subsection, and name
    df = df.sort_values(['Section', 'Subsection', 'Last Name', 'First Name'])
    
    print(f"  Processed {len(df)} employees")
    return df


def create_formatted_output(df, output_path):
    """Create formatted Excel output with multi-column layout."""
    print("Creating formatted Excel output...")
    
    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Staffing Sheet"
    
    # Setup page for 17x11 landscape
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    
    # Set margins
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    
    # Column widths - adjust for better layout
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = 15  # Office subsection header
    for col in range(3, 8):
        ws.column_dimensions[get_column_letter(col)].width = 11
    ws.column_dimensions['H'].width = 2  # Spacer
    for col in range(9, 15):
        ws.column_dimensions[get_column_letter(col)].width = 11
    ws.column_dimensions['O'].width = 2  # Spacer
    for col in range(16, 22):
        ws.column_dimensions[get_column_letter(col)].width = 11
    
    # Current row position
    current_row = 1
    
    # Header styles
    header_font = Font(name='Arial', size=12, bold=True)
    subheader_font = Font(name='Arial', size=10, bold=True)
    data_font = Font(name='Arial', size=10)
    
    # Create section headers
    section_cols = {'Office': 2, 'Fabrication': 9, 'Finishing': 16}
    
    # Section headers (row 1)
    for section, col in section_cols.items():
        cell = ws.cell(row=current_row, column=col)
        cell.value = section
        cell.font = header_font
    
    current_row += 2  # Skip a row
    
    # Process each section
    section_rows = {section: current_row for section in SECTION_ORDER}
    
    for section in SECTION_ORDER:
        col_start = section_cols[section]
        row = section_rows[section]
        
        # Get subsections for this section
        subsections = SUBSECTION_ORDER.get(section, [])
        section_df = df[df['Section'] == section]
        
        for subsection in subsections:
            subsection_df = section_df[section_df['Subsection'] == subsection]
            
            if len(subsection_df) == 0:
                continue
            
            # Check if this subsection uses simple layout
            use_simple_layout = subsection in SIMPLE_LAYOUT_SUBSECTIONS
            
            # Subsection header
            cell = ws.cell(row=row, column=col_start)
            cell.value = subsection
            cell.font = subheader_font
            row += 1
            
            # Column headers
            if use_simple_layout:
                # Simple layout: Name (col+1) and Title (col+4)
                ws.cell(row=row, column=col_start + 1).value = 'Name'
                ws.cell(row=row, column=col_start + 1).font = subheader_font
                ws.cell(row=row, column=col_start + 4).value = 'Title'
                ws.cell(row=row, column=col_start + 4).font = subheader_font
            else:
                # Full layout: Name, Shift, ID #, Title, Codes
                headers = ['Name', 'Shift', 'ID #', 'Title', 'Codes']
                for i, header in enumerate(headers):
                    cell = ws.cell(row=row, column=col_start + i + 1)
                    cell.value = header
                    cell.font = subheader_font
            row += 1
            
            # Employee data
            for _, emp in subsection_df.iterrows():
                name = f"{emp['First Name']} {emp['Last Name']}"
                
                if use_simple_layout:
                    # Simple layout
                    ws.cell(row=row, column=col_start + 1).value = name
                    ws.cell(row=row, column=col_start + 1).font = data_font
                    ws.cell(row=row, column=col_start + 4).value = emp['Simple Title']
                    ws.cell(row=row, column=col_start + 4).font = data_font
                else:
                    # Full layout
                    ws.cell(row=row, column=col_start + 1).value = name
                    ws.cell(row=row, column=col_start + 2).value = emp['Shift']
                    ws.cell(row=row, column=col_start + 3).value = emp['Employee Id']
                    ws.cell(row=row, column=col_start + 4).value = emp['Simple Title']
                    ws.cell(row=row, column=col_start + 5).value = emp['Status Codes']
                    
                    # Apply font to all cells
                    for i in range(5):
                        ws.cell(row=row, column=col_start + i + 1).font = data_font
                
                row += 1
            
            # Add blank row after subsection
            row += 1
        
        # Update section row count
        section_rows[section] = row
    
    # Save workbook
    wb.save(output_path)
    print(f"  Output saved to: {output_path}")


def main():
    """Main function to run the staffing sheet generator."""
    parser = argparse.ArgumentParser(
        description='Generate formatted staffing sheet from HR export files'
    )
    parser.add_argument(
        '--file1',
        default='Emailed - Staffing Emergency Employee List - no grouping.xls',
        help='First input file (default: Emailed - Staffing Emergency Employee List - no grouping.xls)'
    )
    parser.add_argument(
        '--file2',
        default='EmployeeInformation-EmergencyEmployeeList-nogrouping.xlsx',
        help='Second input file (default: EmployeeInformation-EmergencyEmployeeList-nogrouping.xlsx)'
    )
    parser.add_argument(
        '--output',
        help='Output file path (default: Staffing_Sheet_[DATE].xlsx)'
    )
    
    args = parser.parse_args()
    
    # Generate output filename with timestamp if not specified
    if args.output is None:
        timestamp = datetime.now().strftime('%Y-%m-%d')
        args.output = f'Staffing_Sheet_{timestamp}.xlsx'
    
    try:
        print("\n=== Staffing Sheet Generator ===\n")
        
        # Read input files
        df1 = read_employee_file(args.file1)
        df2 = read_employee_file(args.file2)
        
        # Merge data
        merged_df = merge_employee_data(df1, df2)
        
        # Filter to active employees
        active_df = filter_active_employees(merged_df)
        
        # Process employee data
        processed_df = process_employee_data(active_df)
        
        # Create formatted output
        create_formatted_output(processed_df, args.output)
        
        print("\n=== SUCCESS ===")
        print(f"Staffing sheet generated: {args.output}")
        print(f"Total employees: {len(processed_df)}")
        
    except Exception as e:
        print(f"\n=== ERROR ===")
        print(f"Failed to generate staffing sheet: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == '__main__':
    main()
